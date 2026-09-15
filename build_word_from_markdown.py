from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SRC = ROOT / "bao_cao_testcase_bao_mat_burp_project.md"
OUT = ROOT / "bao_cao_testcase_bao_mat_burp_project.docx"


def set_run_font(run, name="Calibri", size=11, bold=False, italic=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color.replace("#", ""))


def set_paragraph_format(paragraph, before=0, after=6, line=1.25, align=None):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if align is not None:
        paragraph.alignment = align


def add_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.allow_autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths)))

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), "120")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(width))
        grid.append(gc)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.tcW
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(widths[idx]))
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_paragraph_with_inline_format(doc, text, style=None, bold=False, italic=False, font_size=11, color=None, align=None):
    p = doc.add_paragraph(style=style)
    set_paragraph_format(p, before=0, after=6, line=1.25, align=align)
    run = p.add_run(text)
    set_run_font(run, size=font_size, bold=bold, italic=italic, color=color)
    return p


def add_markdown_runs(paragraph, text, base_size=11, base_font="Calibri"):
    token_re = re.compile(r"(\*\*.+?\*\*|`.+?`)")
    parts = token_re.split(text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, name=base_font, size=base_size, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, name="Consolas", size=max(base_size - 1, 9))
        else:
            run = paragraph.add_run(part)
            set_run_font(run, name=base_font, size=base_size)


def add_markdown_paragraph(doc, text, style=None, base_size=11, align=None):
    p = doc.add_paragraph(style=style)
    set_paragraph_format(p, before=0, after=6, line=1.25, align=align)
    add_markdown_runs(p, text, base_size=base_size)
    return p


def add_code_block(doc, lines):
    for line in lines:
        p = doc.add_paragraph()
        set_paragraph_format(p, before=0, after=0, line=1.0)
        run = p.add_run(line)
        set_run_font(run, name="Consolas", size=9)
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.right_indent = Inches(0.2)


def table_widths_for(columns, header_texts):
    if columns == 8:
        return [
            650,
            1200,
            850,
            2150,
            1150,
            1600,
            850,
            910,
        ]
    if columns == 2:
        return [2200, 7160]
    if columns == 3:
        return [1800, 3000, 4560]
    if columns == 1:
        return [9360]
    base = 9360 // columns
    widths = [base] * columns
    widths[-1] += 9360 - sum(widths)
    return widths


def add_table(doc, rows):
    columns = len(rows[0])
    widths = table_widths_for(columns, rows[0])
    table = doc.add_table(rows=len(rows), cols=columns)
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    set_repeat_table_header(table.rows[0])

    for r_idx, row in enumerate(rows):
        for c_idx, raw in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx in (0, 6, 7) else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            clean = raw.strip().replace("**", "").replace("`", "")
            run = p.add_run(clean)
            if r_idx == 0:
                set_run_font(run, size=8.5, bold=True)
                add_shading(cell, "E8EEF5")
            else:
                set_run_font(run, size=8)
            if c_idx in (0, 6, 7):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, top=70, start=100, bottom=70, end=100)
    return table


def parse_markdown(md_text: str):
    lines = md_text.splitlines()
    blocks = []
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line.startswith("```"):
            fence = line.strip()
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                code_lines.append(lines[i])
                i += 1
            blocks.append(("code", code_lines))
            i += 1
            continue
        if line.startswith("|"):
            table_lines = [line]
            i += 1
            while i < len(lines) and lines[i].startswith("|"):
                table_lines.append(lines[i])
                i += 1
            rows = []
            for row_line in table_lines:
                parts = [p.strip() for p in row_line.strip().strip("|").split("|")]
                rows.append(parts)
            if len(rows) >= 2 and all(set(cell) <= set("-: ") for cell in rows[1]):
                rows.pop(1)
            blocks.append(("table", rows))
            continue
        if line.startswith("#"):
            hashes = len(line) - len(line.lstrip("#"))
            blocks.append(("heading", hashes, line[hashes:].strip()))
            i += 1
            continue
        if line.startswith("- "):
            items = [line[2:].strip()]
            i += 1
            while i < len(lines) and lines[i].startswith("- "):
                items.append(lines[i][2:].strip())
                i += 1
            blocks.append(("bullets", items))
            continue
        if line.strip() == "---":
            blocks.append(("rule",))
            i += 1
            continue
        blocks.append(("para", line.strip()))
        i += 1
    return blocks


def build_doc():
    md_text = SRC.read_text(encoding="utf-8")
    blocks = parse_markdown(md_text)
    doc = Document()

    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 18, 10),
        ("Heading 2", 13, "2E74B5", 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ]:
        st = styles[style_name]
        st.font.name = "Calibri"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)

    # Title style is custom, not the built-in Word Title style.
    if "DocTitle" not in styles:
        title_style = styles.add_style("DocTitle", 1)
        title_style.font.name = "Calibri"
        title_style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        title_style.font.size = Pt(18)
        title_style.font.bold = True
        title_style.font.color.rgb = RGBColor.from_string("1F4D78")
        title_style.paragraph_format.space_before = Pt(0)
        title_style.paragraph_format.space_after = Pt(10)

    first_heading_done = False
    for block in blocks:
        kind = block[0]
        if kind == "heading":
            level, text = block[1], block[2]
            if level == 1 and not first_heading_done:
                p = doc.add_paragraph(style="DocTitle")
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                add_markdown_runs(p, text, base_size=18)
                for run in p.runs:
                    run.bold = True
                    run.font.size = Pt(18)
                first_heading_done = True
            else:
                p = doc.add_paragraph(style=f"Heading {min(level,3)}")
                add_markdown_runs(p, text, base_size={2: 16, 3: 13}.get(level, 12))
                if level == 2:
                    p.style = styles["Heading 1"]
                elif level == 3:
                    p.style = styles["Heading 2"]
                elif level >= 4:
                    p.style = styles["Heading 3"]
            continue
        if kind == "para":
            text = block[1]
            if not text:
                continue
            if text.startswith("**") and text.endswith("**"):
                p = doc.add_paragraph()
                set_paragraph_format(p, before=0, after=6, line=1.25)
                run = p.add_run(text.strip("*"))
                set_run_font(run, size=11, bold=True)
            else:
                add_markdown_paragraph(doc, text)
            continue
        if kind == "bullets":
            for item in block[1]:
                p = doc.add_paragraph(style="List Bullet")
                set_paragraph_format(p, before=0, after=4, line=1.25)
                add_markdown_runs(p, item, base_size=11)
            continue
        if kind == "rule":
            p = doc.add_paragraph()
            set_paragraph_format(p, before=0, after=6, line=1.0)
            p.add_run("")
            continue
        if kind == "code":
            add_code_block(doc, block[1])
            continue
        if kind == "table":
            table = add_table(doc, block[1])
            # Add a little spacing after tables.
            spacer = doc.add_paragraph()
            set_paragraph_format(spacer, before=0, after=6, line=1.0)
            continue

    doc.save(OUT)


if __name__ == "__main__":
    build_doc()
    print(str(OUT))
